import os
import re
import time
from datetime import datetime
from pathlib import Path

import pyautogui
import pytesseract
from PIL import ImageEnhance, ImageOps
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


# Change this if Tesseract is installed in a different folder.
pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)

pyautogui.FAILSAFE = True
pyautogui.PAUSE = 1

SOURCE_URL = "https://www.grtjewels.com"
OUTPUT_FOLDER = Path(__file__).resolve().parent


def open_grt_in_edge():
    """Open Microsoft Edge and navigate to the GRT Jewellers website."""
    pyautogui.hotkey("win", "r")
    time.sleep(1)
    pyautogui.write("msedge")
    pyautogui.press("enter")
    time.sleep(4)

    pyautogui.hotkey("ctrl", "l")
    pyautogui.write(SOURCE_URL, interval=0.03)
    pyautogui.press("enter")
    time.sleep(15)

    # Maximize Edge so the rate appears in the expected screen area.
    pyautogui.hotkey("win", "up")
    time.sleep(2)


def read_gold_rate_from_screen():
    """Capture the screen and use OCR to read the 22KT, 1 g gold rate."""
    screen_image = pyautogui.screenshot()

    # Save a screenshot of the GRT page in the script's folder.
    browser_screenshot = OUTPUT_FOLDER / "screenshot.png"
    screen_image.save(browser_screenshot)
    print(f"Browser screenshot saved: {browser_screenshot}")

    screen_width, screen_height = screen_image.size

    # Crop the screen area where the gold-rate text appears.
    rate_area = screen_image.crop((
        int(screen_width * 0.57),
        int(screen_height * 0.08),
        int(screen_width * 0.77),
        int(screen_height * 0.15),
    ))

    # Enlarge and improve contrast to help OCR read the text.
    rate_area = ImageOps.grayscale(rate_area)
    rate_area = rate_area.resize(
        (rate_area.width * 4, rate_area.height * 4)
    )
    rate_area = ImageEnhance.Contrast(rate_area).enhance(2)

    recognized_text = pytesseract.image_to_string(
        rate_area,
        config="--psm 6"
    )
    print("Text recognized from rate area:", recognized_text.strip())

    # The displayed text is similar to: GOLD 22 KT/1g - ₹13675
    match = re.search(
        r"GOLD\s*22\s*KT.*?(\d[\d,]{3,}(?:\.\d{1,2})?)",
        recognized_text,
        flags=re.IGNORECASE | re.DOTALL,
    )

    if not match:
        raise ValueError(
            "Could not read the 22KT gold rate from the screenshot. "
            "Check that the rate is visible near the top of the page."
        )

    return float(match.group(1).replace(",", ""))


def create_excel_report(gold_rate, run_time):
    """Create the dated Excel report in the script's folder."""
    date_for_filename = run_time.strftime("%Y-%m-%d")
    excel_path = OUTPUT_FOLDER / f"daily_report_{date_for_filename}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Gold_Price"

    sheet.append(["Date", "Time", "Gold Rate (22KT 1 g)", "Source"])
    sheet.append([
        run_time.date(),
        run_time.time().replace(microsecond=0),
        gold_rate,
        SOURCE_URL,
    ])

    # Format the header row.
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    sheet["A2"].number_format = "yyyy-mm-dd"
    sheet["B2"].number_format = "hh:mm:ss"
    sheet["C2"].number_format = '"₹"#,##0.00'

    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 28
    sheet.column_dimensions["D"].width = 35

    workbook.save(excel_path)
    return excel_path


def save_excel_screenshot(excel_path, run_time):
    """Open the report in Excel and save a screenshot of the sheet."""
    os.startfile(str(excel_path))
    time.sleep(5)

    pyautogui.hotkey("win", "up")
    pyautogui.hotkey("ctrl", "home")
    time.sleep(1)

    screenshot_path = OUTPUT_FOLDER / (
        f"daily_report_{run_time:%Y-%m-%d}.png"
    )
    pyautogui.screenshot().save(screenshot_path)
    return screenshot_path


def main():
    try:
        print("Opening GRT Jewellers in Microsoft Edge...")
        open_grt_in_edge()

        gold_rate = read_gold_rate_from_screen()
        run_time = datetime.now()

        excel_path = create_excel_report(gold_rate, run_time)
        print(f"22KT gold rate captured: ₹{gold_rate:,.2f}")
        print(f"Excel report saved: {excel_path}")

        screenshot_path = save_excel_screenshot(excel_path, run_time)
        print(f"Excel screenshot saved: {screenshot_path}")

    except Exception as error:
        print(f"ERROR TYPE: {type(error).__name__}")
        print(f"DETAILS: {error}")


if __name__ == "__main__":
    main()